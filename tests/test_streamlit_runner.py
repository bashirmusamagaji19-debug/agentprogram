from __future__ import annotations

import base64
import json
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from web_task_agent.streamlit_runner import (
    UiRequestError,
    UiRunRequest,
    decode_resume_upload,
    extract_resume_upload,
    parse_seed_urls,
    parse_skills,
    read_download_artifact,
    run_ui_request,
    validate_ui_request,
)

_PDF_RESUME = base64.b64decode(
    """
    JVBERi0xLjMKJZOMi54gUmVwb3J0TGFiIEdlbmVyYXRlZCBQREYgZG9jdW1lbnQgKG9wZW5zb3Vy
    Y2UpCjEgMCBvYmoKPDwKL0YxIDIgMCBSCj4+CmVuZG9iagoyIDAgb2JqCjw8Ci9CYXNlRm9udCAv
    SGVsdmV0aWNhIC9FbmNvZGluZyAvV2luQW5zaUVuY29kaW5nIC9OYW1lIC9GMSAvU3VidHlwZSAv
    VHlwZTEgL1R5cGUgL0ZvbnQKPj4KZW5kb2JqCjMgMCBvYmoKPDwKL0NvbnRlbnRzIDcgMCBSIC9N
    ZWRpYUJveCBbIDAgMCA1OTUuMjc1NiA4NDEuODg5OCBdIC9QYXJlbnQgNiAwIFIgL1Jlc291cmNl
    cyA8PAovRm9udCAxIDAgUiAvUHJvY1NldCBbIC9QREYgL1RleHQgL0ltYWdlQiAvSW1hZ2VDIC9J
    bWFnZUkgXQo+PiAvUm90YXRlIDAgL1RyYW5zIDw8Cgo+PiAKICAvVHlwZSAvUGFnZQo+PgplbmRv
    YmoKNCAwIG9iago8PAovUGFnZU1vZGUgL1VzZU5vbmUgL1BhZ2VzIDYgMCBSIC9UeXBlIC9DYXRh
    bG9nCj4+CmVuZG9iago1IDAgb2JqCjw8Ci9BdXRob3IgKGFub255bW91cykgL0NyZWF0aW9uRGF0
    ZSAoRDoyMDI2MDkyMTEzMTE0MSswOCcwMCcpIC9DcmVhdG9yIChhbm9ueW1vdXMpIC9LZXl3b3Jk
    cyAoKSAvTW9kRGF0ZSAoRDoyMDI2MDkyMTEzMTE0MSswOCcwMCcpIC9Qcm9kdWNlciAoUmVwb3J0
    TGFiIFBERiBMaWJyYXJ5IC0gXChvcGVuc291cmNlXCkpIAogIC9TdWJqZWN0ICh1bnNwZWNpZmll
    ZCkgL1RpdGxlICh1bnRpdGxlZCkgL1RyYXBwZWQgL0ZhbHNlCj4+CmVuZG9iago2IDAgb2JqCjw8
    Ci9Db3VudCAxIC9LaWRzIFsgMyAwIFIgXSAvVHlwZSAvUGFnZXMKPj4KZW5kb2JqCjcgMCBvYmoK
    PDwKL0ZpbHRlciBbIC9BU0NJSTg1RGVjb2RlIC9GbGF0ZURlY29kZSBdIC9MZW5ndGggMTYyCj4+
    CnN0cmVhbQpHYXBRaDBFPUYsMFVcSDNUXHBOWVReUUtrP3RjPklQLDtXI1UxXjIzaWhQRU1fP0NX
    NEtJU2k8IVs3YCNPQl9zS2k5XFtYUzc0IiZKQkYlNnImYnUnVD5cQmFwUyxUWCRsdCRFN3NAW11P
    PVoocy9NLFRiK1BIcFRvITs1WDk8OGdJNFZcc2M7dEVUcTRTKTBDTXQycmQuVz5DZydwMG1nfj5l
    bmRzdHJlYW0KZW5kb2JqCnhyZWYKMCA4CjAwMDAwMDAwMDAgNjU1MzUgZiAKMDAwMDAwMDA2MSAw
    MDAwMCBuIAowMDAwMDAwMDkyIDAwMDAwIG4gCjAwMDAwMDAxOTkgMDAwMDAgbiAKMDAwMDAwMDQw
    MiAwMDAwMCBuIAowMDAwMDAwNDcwIDAwMDAwIG4gCjAwMDAwMDA3MzEgMDAwMDAgbiAKMDAwMDAw
    MDc5MCAwMDAwMCBuIAp0cmFpbGVyCjw8Ci9JRCAKWzw3MDY2OTBkNWYxYTA5MmEzYjczMWM2NjA4
    ODQ5NTcxYj48NzA2NjkwZDVmMWEwOTJhM2I3MzFjNjYwODg0OTU3MWI+XQolIFJlcG9ydExhYiBn
    ZW5lcmF0ZWQgUERGIGRvY3VtZW50IC0tIGRpZ2VzdCAob3BlbnNvdXJjZSkKCi9JbmZvIDUgMCBS
    Ci9Sb290IDQgMCBSCi9TaXplIDgKPj4Kc3RhcnR4cmVmCjEwNDIKJSVFT0YK
    """
)


def test_parse_skills_normalizes_commas_lines_and_duplicates() -> None:
    assert parse_skills("Python，RAG\npython, LangGraph") == [
        "Python",
        "RAG",
        "LangGraph",
    ]


def test_parse_seed_urls_keeps_unique_http_urls() -> None:
    assert parse_seed_urls(
        "https://example.com/jobs/1\nhttps://example.com/jobs/1\nhttp://example.com/jobs/2"
    ) == ["https://example.com/jobs/1", "http://example.com/jobs/2"]


def test_decode_resume_upload_accepts_utf8_text() -> None:
    assert decode_resume_upload("项目：LangGraph Agent".encode()) == "项目：LangGraph Agent"


def test_decode_resume_upload_rejects_invalid_utf8() -> None:
    with pytest.raises(UiRequestError, match="UTF-8"):
        decode_resume_upload(b"\xff\xfe")


def test_extract_resume_upload_reads_pdf_text() -> None:
    resume = extract_resume_upload(_PDF_RESUME, filename="resume.pdf")

    assert "Python" in resume
    assert "LangGraph" in resume
    assert "RAG" in resume


def test_extract_resume_upload_reads_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("项目经历")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技术栈"
    table.cell(0, 1).text = "Python / LangGraph"
    buffer = BytesIO()
    document.save(buffer)

    resume = extract_resume_upload(buffer.getvalue(), filename="resume.docx")

    assert "项目经历" in resume
    assert "技术栈" in resume
    assert "Python / LangGraph" in resume


def test_extract_resume_upload_rejects_pdf_without_text() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(UiRequestError, match="文本|扫描"):
        extract_resume_upload(buffer.getvalue(), filename="scan.pdf")


def test_seed_url_mode_requires_http_urls() -> None:
    request = UiRunRequest(data_mode="seed_urls", seed_urls=["not-a-url"])

    with pytest.raises(UiRequestError, match="HTTP"):
        validate_ui_request(request, environ={})


def test_aggregator_mode_requires_uploaded_json_path() -> None:
    request = UiRunRequest(data_mode="aggregator")

    with pytest.raises(UiRequestError, match="聚合岗位 JSON"):
        validate_ui_request(request, environ={})


def test_provider_validation_names_missing_variable_without_value() -> None:
    request = UiRunRequest(llm_extractor_provider="qwen")

    with pytest.raises(UiRequestError, match="DASHSCOPE_API_KEY") as captured:
        validate_ui_request(request, environ={})

    assert "secret-value" not in str(captured.value)


def test_provider_validation_accepts_configured_environment() -> None:
    request = UiRunRequest(
        llm_extractor_provider="deepseek",
        llm_match_provider="qwen",
    )

    validate_ui_request(
        request,
        environ={"DEEPSEEK_API_KEY": "secret-value", "DASHSCOPE_API_KEY": "secret-value"},
    )


def test_request_rejects_target_count_above_ui_limit() -> None:
    with pytest.raises(ValueError):
        UiRunRequest(target_count=51)


@pytest.mark.asyncio
async def test_aggregator_request_passes_keyword_to_source(monkeypatch, tmp_path) -> None:
    import web_task_agent.streamlit_runner as runner_module
    from web_task_agent.models import BrowserPage

    captured: dict[str, str | int] = {}

    async def fake_discover(self, limit: int, *, keyword: str = ""):
        captured["limit"] = limit
        captured["keyword"] = keyword
        return []

    monkeypatch.setattr(runner_module.AggregatorRepoSource, "discover", fake_discover)

    async def empty_page_loader(url: str):
        return BrowserPage(url=url, title="", content="")

    await run_ui_request(
        UiRunRequest(
            data_mode="aggregator",
            aggregator_url="https://example.com/jobs.json",
            keyword="AI Agent 实习",
            target_count=7,
        ),
        output_root=tmp_path,
        environ={},
        page_loader=empty_page_loader,
    )

    assert captured == {"limit": 7, "keyword": "AI Agent 实习"}


@pytest.mark.asyncio
async def test_official_source_diagnostics_and_page_diagnostics_are_in_result_and_json(
    monkeypatch,
    tmp_path,
) -> None:
    import web_task_agent.streamlit_runner as runner_module
    from web_task_agent.browser import PageHttpError
    from web_task_agent.job_sources import DiscoveredJob

    captured: dict[str, str | int] = {}

    class FakeOfficialSource:
        def __init__(self, specs=None):  # noqa: ANN001
            self.diagnostics: list[dict[str, object]] = []

        async def discover(self, limit: int, *, keyword: str = ""):
            captured["limit"] = limit
            captured["keyword"] = keyword
            self.diagnostics = [
                {
                    "source": "broken",
                    "status": "failed",
                    "attempts": 1,
                    "jobs_discovered": 0,
                    "error_type": "RuntimeError",
                    "detail": "simulated failure",
                }
            ]
            return [
                DiscoveredJob(
                    url="https://example.com/jobs/1",
                    title="大模型算法实习生",
                    company="示例公司",
                    location="北京",
                    jd_text=(
                        "岗位职责：负责大模型训练、评测和推理优化。"
                        "任职要求：熟悉 Python、PyTorch 和 Transformer。"
                    )
                    * 4,
                    source="healthy",
                )
            ]

    async def blocked_page_loader(url: str):
        raise PageHttpError(f"HTTP 403: {url}", status_code=403)

    monkeypatch.setattr(runner_module, "OfficialListSource", FakeOfficialSource)

    result = await run_ui_request(
        UiRunRequest(
            data_mode="official",
            official_specs=["broken"],
            keyword="大模型实习",
            target_count=3,
        ),
        output_root=tmp_path,
        environ={},
        page_loader=blocked_page_loader,
    )

    assert captured == {"limit": 3, "keyword": "大模型实习"}
    assert result.jobs[0].url_verification == "list-attested"
    assert result.diagnostics[0] == {
        "url": "broken",
        "error": "source_failed: RuntimeError: simulated failure",
    }
    assert result.diagnostics[1]["url"] == "https://example.com/jobs/1"
    assert result.diagnostics[1]["error"].startswith(
        "page_blocked: http-failed:PageHttpError: HTTP 403"
    )

    payload = json.loads(result.artifacts["json"].read_text(encoding="utf-8"))
    assert set(payload["metadata"]) == {"source_diagnostics", "page_diagnostics"}
    assert payload["metadata"]["source_diagnostics"][0]["source"] == "broken"
    assert payload["metadata"]["page_diagnostics"][0]["category"] == "blocked"


def test_source_status_rows_distinguish_empty_and_unknown() -> None:
    import web_task_agent.streamlit_runner as runner_module

    rows = runner_module._source_diagnostic_rows(
        [
            {"source": "empty-source", "status": "empty"},
            {"source": "unknown-source", "status": "unknown"},
            {"source": "missing-status-source"},
        ]
    )

    assert [row["error"] for row in rows] == [
        "source_empty: no jobs returned",
        "source_unknown: no jobs returned",
        "source_unknown: no jobs returned",
    ]


@pytest.mark.asyncio
async def test_blocked_page_preserves_terminal_workflow_failure(
    monkeypatch,
    tmp_path,
) -> None:
    import web_task_agent.streamlit_runner as runner_module
    from web_task_agent.browser import PageHttpError
    from web_task_agent.job_sources import DiscoveredJob

    url = "https://example.com/jobs/blocked-short-jd"

    class FakeOfficialSource:
        def __init__(self, specs=None):  # noqa: ANN001
            self.diagnostics: list[dict[str, object]] = []

        async def discover(self, limit: int, *, keyword: str = ""):
            return [
                DiscoveredJob(
                    url=url,
                    title="大模型算法实习生",
                    company="示例公司",
                    location="北京",
                    jd_text="岗位职责：模型训练。",
                    source="official-test",
                )
            ]

    async def blocked_page_loader(_url: str):
        raise PageHttpError("HTTP 403 forbidden", status_code=403)

    monkeypatch.setattr(runner_module, "OfficialListSource", FakeOfficialSource)

    result = await run_ui_request(
        UiRunRequest(
            data_mode="official",
            official_specs=["test"],
            target_count=1,
        ),
        output_root=tmp_path,
        environ={},
        page_loader=blocked_page_loader,
    )

    assert result.jobs == []
    assert result.failed_urls == [url]
    assert len(result.diagnostics) == 2
    assert result.diagnostics[0]["url"] == url
    assert result.diagnostics[0]["error"].startswith("page_blocked:")
    assert result.diagnostics[1]["url"] == url
    assert result.diagnostics[1]["error"].startswith("PageEmptyError:")


@pytest.mark.asyncio
async def test_demo_mode_has_no_official_source_diagnostics(tmp_path) -> None:
    result = await run_ui_request(
        UiRunRequest(data_mode="demo"),
        output_root=tmp_path,
        environ={},
    )

    payload = json.loads(result.artifacts["json"].read_text(encoding="utf-8"))

    assert payload["metadata"] == {}
    assert "source_diagnostics" not in payload["metadata"]
    assert all(not item["error"].startswith("source_") for item in result.diagnostics)


@pytest.mark.asyncio
async def test_page_diagnostic_deduplicates_failed_url_and_filtered_row_stays_last(
    monkeypatch,
    tmp_path,
) -> None:
    import web_task_agent.streamlit_runner as runner_module
    from web_task_agent.browser import BrowserPage, PageHttpError
    from web_task_agent.job_sources import DiscoveredJob

    removed_url = "https://example.com/jobs/removed"
    filtered_url = "https://example.com/jobs/other-location"

    class FakeOfficialSource:
        def __init__(self, specs=None):  # noqa: ANN001
            self.diagnostics: list[dict[str, object]] = []

        async def discover(self, limit: int, *, keyword: str = ""):
            return [
                DiscoveredJob(
                    url=removed_url,
                    title="大模型算法实习生",
                    company="示例公司",
                    location="北京",
                    jd_text="岗位职责：模型训练。" * 20,
                    source="official-test",
                ),
                DiscoveredJob(
                    url=filtered_url,
                    title="大模型算法实习生",
                    company="示例公司",
                    location="上海",
                    jd_text="岗位职责：模型训练。" * 20,
                    source="official-test",
                ),
            ]

    async def mixed_page_loader(url: str):
        if url == removed_url:
            raise PageHttpError("HTTP 404 gone", status_code=404)
        return BrowserPage(
            url=url,
            title="大模型算法实习生",
            content=(
                "岗位名称：大模型算法实习生\n"
                "公司：示例公司\n"
                "工作地点：上海\n"
                "岗位职责：负责大模型训练、评测和推理优化，提升线上系统效果。\n"
                "任职要求：熟悉 Python、PyTorch 和 Transformer，具备扎实的工程能力。"
                "能够持续实习并协作完成复杂项目。"
            ),
        )

    monkeypatch.setattr(runner_module, "OfficialListSource", FakeOfficialSource)

    result = await run_ui_request(
        UiRunRequest(
            data_mode="official",
            official_specs=["test"],
            location="北京",
            target_count=2,
        ),
        output_root=tmp_path,
        environ={},
        page_loader=mixed_page_loader,
    )

    assert result.failed_urls == [removed_url]
    assert len(result.diagnostics) == 2
    assert result.diagnostics[0]["url"] == removed_url
    assert result.diagnostics[0]["error"].startswith("page_removed:")
    assert result.diagnostics[-1]["error"].startswith("verification_filtered:")


@pytest.mark.asyncio
async def test_demo_request_returns_jobs_diagnostics_and_artifacts(tmp_path) -> None:
    result = await run_ui_request(
        UiRunRequest(
            keyword="AI intern",
            location="全国",
            target_count=2,
            skills=["Python", "LangGraph"],
            resume_text="Built a LangGraph job agent with Python.",
        ),
        output_root=tmp_path,
        environ={},
    )

    assert len(result.jobs) == 2
    assert result.metrics.valid_jobs == 2
    assert result.execution_trace
    assert set(result.artifacts) == {"json", "report", "dashboard", "action_plan"}
    assert all(path.is_file() for path in result.artifacts.values())
    assert result.artifacts["json"].parent.name == result.run_id


@pytest.mark.asyncio
async def test_each_ui_run_uses_an_isolated_output_directory(tmp_path) -> None:
    first = await run_ui_request(UiRunRequest(), output_root=tmp_path, environ={})
    second = await run_ui_request(UiRunRequest(), output_root=tmp_path, environ={})

    assert first.run_id != second.run_id
    assert first.artifacts["json"].parent != second.artifacts["json"].parent


@pytest.mark.asyncio
async def test_seed_url_failure_is_returned_as_diagnostic(tmp_path) -> None:
    result = await run_ui_request(
        UiRunRequest(
            data_mode="seed_urls",
            seed_urls=["https://example.invalid/jobs/missing"],
        ),
        output_root=tmp_path,
        environ={},
        page_loader=lambda _url: (_ for _ in ()).throw(ValueError("page unavailable")),
    )

    assert result.jobs == []
    assert result.failed_urls == ["https://example.invalid/jobs/missing"]
    assert "page unavailable" in result.diagnostics[0]["error"]


@pytest.mark.asyncio
async def test_download_reader_only_accepts_registered_artifacts(tmp_path) -> None:
    result = await run_ui_request(UiRunRequest(), output_root=tmp_path, environ={})
    unrelated = tmp_path / "secret.txt"
    unrelated.write_text("do not expose", encoding="utf-8")

    name, content = read_download_artifact(result, "report")
    assert name.endswith(".md")
    assert "AI 实习岗位搜索报告" in content.decode("utf-8")

    with pytest.raises(UiRequestError, match="不可下载"):
        read_download_artifact(result, "secret")
